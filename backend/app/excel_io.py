"""Excel import and export.

Import: reads an existing tracker spreadsheet (like the user's sample) and
maps its columns into Jobs + Applications, normalizing dates/statuses and
detecting duplicates. The credential column is intentionally ignored -- we
never store portal passwords.

Export: produces a multi-sheet workbook (Applications, Contacts, Follow-ups,
Analytics) from the database.
"""
from __future__ import annotations

import io
from datetime import date

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session, selectinload

from . import crud, models, schemas
from .utils import parse_date, parse_sponsorship, parse_status

# Column header aliases -> canonical field. Matching is case-insensitive and
# ignores surrounding whitespace so slightly different templates still work.
_HEADER_MAP = {
    "internship/job company": "company",
    "company": "company",
    "position": "position",
    "job title": "position",
    "role": "position",
    "description": "description",
    "status": "status",
    "link": "link",
    "job link": "link",
    "job url": "link",
    "visa sponsorship(yes/no)": "sponsorship",
    "visa sponsorship (yes/no)": "sponsorship",
    "visa sponsorship": "sponsorship",
    "sponsorship": "sponsorship",
    "contact info": "contact",
    "contact": "contact",
    "date month": "date_applied",
    "date applied": "date_applied",
    "applied date": "date_applied",
    # "Username and pass" and "Ageing" are intentionally NOT mapped:
    # credentials are never stored; ageing is always computed.
}


def _norm_header(value) -> str:
    return str(value or "").strip().lower()


class ImportResult:
    def __init__(self) -> None:
        self.imported = 0
        self.duplicates = 0
        self.skipped = 0
        self.credential_columns_ignored = False
        self.warnings: list[str] = []

    def as_dict(self) -> dict:
        return {
            "imported": self.imported,
            "duplicates": self.duplicates,
            "skipped": self.skipped,
            "credential_columns_ignored": self.credential_columns_ignored,
            "warnings": self.warnings,
        }


def import_workbook(db: Session, content: bytes) -> ImportResult:
    result = ImportResult()
    wb = load_workbook(io.BytesIO(content), data_only=True)
    ws = wb.active

    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        result.warnings.append("Spreadsheet is empty.")
        return result

    header_row = rows[0]
    col_index: dict[str, int] = {}
    for idx, cell in enumerate(header_row):
        h = _norm_header(cell)
        if "username" in h or "pass" in h:
            result.credential_columns_ignored = True
            continue
        field = _HEADER_MAP.get(h)
        if field:
            col_index[field] = idx

    if "company" not in col_index and "position" not in col_index:
        result.warnings.append(
            "Could not find a Company or Position column. Check the headers."
        )
        return result

    def cell(row, field):
        i = col_index.get(field)
        if i is None or i >= len(row):
            return None
        return row[i]

    for row in rows[1:]:
        if row is None or all(c is None or str(c).strip() == "" for c in row):
            continue
        company = cell(row, "company")
        position = cell(row, "position")
        if not company and not position:
            result.skipped += 1
            continue

        data = schemas.ApplicationCreate(
            company_name=str(company).strip() if company else None,
            position=str(position).strip() if position else "Untitled Position",
            description=(str(cell(row, "description")).strip()
                        if cell(row, "description") else None),
            link=str(cell(row, "link")).strip() if cell(row, "link") else None,
            status=parse_status(cell(row, "status")),
            sponsorship_status=parse_sponsorship(cell(row, "sponsorship")),
            date_applied=parse_date(cell(row, "date_applied")),
            notes=(f"Contact: {cell(row, 'contact')}"
                   if cell(row, "contact") else None),
        )

        # Duplicate detection: same company + same position already applied.
        company_obj = crud.get_or_create_company(db, data.company_name)
        company_id = company_obj.id if company_obj else None
        existing_job = crud.find_duplicate_job(db, company_id, data.position or "")
        if existing_job is not None and existing_job.application is not None:
            result.duplicates += 1
            continue

        crud.create_application(db, data)
        result.imported += 1

    db.commit()
    return result


# --- Export ----------------------------------------------------------------

_HEADER_FILL = PatternFill("solid", fgColor="1F2937")
_HEADER_FONT = Font(color="FFFFFF", bold=True)


def _write_sheet(ws, headers: list[str], rows: list[list]) -> None:
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
    for r in rows:
        ws.append(r)
    # Auto-ish width.
    for c in range(1, len(headers) + 1):
        letter = get_column_letter(c)
        max_len = len(str(headers[c - 1]))
        for r in rows:
            if c - 1 < len(r) and r[c - 1] is not None:
                max_len = max(max_len, len(str(r[c - 1])))
        ws.column_dimensions[letter].width = min(max_len + 2, 60)
    ws.freeze_panes = "A2"


def export_workbook(db: Session) -> bytes:
    wb = Workbook()

    # Sheet 1: Applications (mirrors the user's tracker + computed ageing).
    ws_apps = wb.active
    ws_apps.title = "Applications"
    apps = (
        db.execute(
            select(models.Application).options(
                selectinload(models.Application.job).selectinload(models.Job.company),
                selectinload(models.Application.resume_version),
            )
        )
        .scalars()
        .all()
    )
    app_rows = []
    for a in apps:
        job = a.job
        company = job.company.name if job and job.company else ""
        app_rows.append([
            company,
            job.title if job else "",
            a.status.value,
            job.job_url if job else "",
            job.sponsorship_status.value if job else "",
            a.resume_version.label if a.resume_version else "",
            a.ats_match,
            a.missing_keywords,
            a.date_applied.isoformat() if a.date_applied else "",
            a.ageing if a.ageing is not None else "",
            a.response or "",
            a.notes or "",
        ])
    _write_sheet(
        ws_apps,
        ["Company", "Position", "Status", "Link", "Visa Sponsorship",
         "Resume Version", "ATS Match", "Missing Keywords", "Date Applied",
         "Ageing", "Response", "Notes"],
        app_rows,
    )

    # Sheet 2: Contacts.
    ws_contacts = wb.create_sheet("Contacts")
    contacts = (
        db.execute(select(models.Contact).options(selectinload(models.Contact.company)))
        .scalars()
        .all()
    )
    contact_rows = [
        [
            c.name,
            c.title or "",
            c.company.name if c.company else "",
            c.email or "",
            "Yes" if c.email_verified else "No",
            c.linkedin_url or "",
            c.contact_type.value,
            c.date_contacted.isoformat() if c.date_contacted else "",
            c.response or "",
        ]
        for c in contacts
    ]
    _write_sheet(
        ws_contacts,
        ["Name", "Role", "Company", "Email", "Email Verified", "LinkedIn",
         "Type", "Date Contacted", "Response"],
        contact_rows,
    )

    # Sheet 3: Follow-ups.
    ws_fu = wb.create_sheet("Follow-ups")
    fus = (
        db.execute(
            select(models.FollowUp).options(
                selectinload(models.FollowUp.application)
                .selectinload(models.Application.job)
                .selectinload(models.Job.company)
            )
        )
        .scalars()
        .all()
    )
    fu_rows = []
    for f in fus:
        app = f.application
        job = app.job if app else None
        company = job.company.name if job and job.company else ""
        fu_rows.append([
            company,
            job.title if job else "",
            f.due_date.isoformat(),
            f.status.value,
            f.note or "",
        ])
    _write_sheet(
        ws_fu,
        ["Company", "Position", "Follow-up Due", "Status", "Note"],
        fu_rows,
    )

    # Sheet 4: Analytics.
    ws_an = wb.create_sheet("Analytics")
    stats = crud.compute_dashboard(db)
    total = stats.total_applications or 1  # avoid div-by-zero
    interview_rate = round(stats.interviews / total * 100, 1)
    responded = stats.interviews + stats.offers + stats.rejections
    response_rate = round(responded / total * 100, 1)
    analytics_rows = [
        ["Total Applications", stats.total_applications],
        ["Applied This Week", stats.applied_this_week],
        ["Interviews", stats.interviews],
        ["Offers", stats.offers],
        ["Rejections", stats.rejections],
        ["Pending Response", stats.pending_response],
        ["Follow-ups Due", stats.follow_ups_due],
        ["Saved (not applied)", stats.jobs_saved_not_applied],
        ["Interview Rate %", interview_rate],
        ["Response Rate %", response_rate],
        ["Generated", date.today().isoformat()],
    ]
    _write_sheet(ws_an, ["Metric", "Value"], analytics_rows)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
